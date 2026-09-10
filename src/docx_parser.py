import base64
import zipfile

from docx import Document

from .models import DocumentData, PageData, TableData, ImageData


def extract_docx(docx_path: str) -> DocumentData:

    # Create our main Pydantic document object
    document = DocumentData(
        source_file=docx_path,
        file_type="docx"
    )

    # For now, we treat the DOCX content as one logical page.
    page_data = PageData(
        page_number=1,
        text=""
    )

    # -----------------------------------
    # Open DOCX
    # -----------------------------------
    doc = Document(docx_path)

    # -----------------------------------
    # Extract paragraphs
    # -----------------------------------
    paragraphs = []

    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    page_data.text = "\n".join(paragraphs)

    # -----------------------------------
    # Extract tables
    # -----------------------------------
    for table in doc.tables:

        markdown_table = convert_table_to_markdown(
            table
        )

        if markdown_table:

            page_data.tables.append(
                TableData(
                    page_number=1,
                    content=markdown_table
                )
            )

    # -----------------------------------
    # Extract embedded images
    # -----------------------------------
    with zipfile.ZipFile(docx_path, "r") as docx_zip:

        for file_name in docx_zip.namelist():

            # Images inside DOCX are stored
            # inside the word/media folder.
            if file_name.startswith("word/media/"):

                image_bytes = docx_zip.read(file_name)

                image_base64 = base64.b64encode(
                    image_bytes
                ).decode("utf-8")

                image_format = file_name.split(".")[-1]

                page_data.images.append(
                    ImageData(
                        page_number=1,
                        image_base64=image_base64,
                        image_format=image_format
                    )
                )

    # Add the processed page to the document
    document.pages.append(page_data)

    return document


def convert_table_to_markdown(table):

    if not table.rows:
        return ""

    rows = []

    # Extract every cell from every row
    for row in table.rows:

        cells = []

        for cell in row.cells:

            text = cell.text.strip()

            cells.append(text)

        rows.append(cells)

    if not rows:
        return ""

    # First row becomes the header
    header = rows[0]

    markdown = "| " + " | ".join(header) + " |\n"

    # Markdown separator
    markdown += "| " + " | ".join(
        ["---"] * len(header)
    ) + " |\n"

    # Remaining rows
    for row in rows[1:]:

        markdown += "| " + " | ".join(row) + " |\n"

    return markdown